"""Synthetic parser fixtures only: no replacement regulatory documents are generated."""
import csv
import json
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pydantic import ValidationError
from hydroguard.ingestion.pipeline import Metadata, extract, ingest, load_passages, retrieve_passages, split_unit
from hydroguard.ingestion.zetaris import registration_plan, query_registered_passage
from scripts.zetaris_smoke import Config


def make_pdf(path, text="Water baseline observations for a synthetic parser test only."):
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 60), text)
    pdf.new_page()  # empty page must not yield invented text
    pdf.save(path)
    pdf.close()


def test_pdf_pages_and_review_flags(tmp_path):
    path = tmp_path / "parser.pdf"
    make_pdf(path)
    units, flags = extract(path)
    assert len(units) == 1 and units[0]["page"] == 1
    assert units[0]["locator"] == "pdf-page:1"
    assert flags[0]["page"] == 2
    assert "ocr_or_manual_review" in flags[0]["reason"]


def test_image_only_pdf_flagged_without_passages(tmp_path):
    path = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    # Image-only test page made from a rendered blank page, not a source document.
    bitmap = page.get_pixmap().tobytes("png")
    page.insert_image(page.rect, stream=bitmap)
    pdf.save(path)
    pdf.close()
    units, flags = extract(path)
    assert units == [] and flags[0]["page"] == 1


def test_docx_headings_paragraphs_and_tables_without_pages(tmp_path):
    path = tmp_path / "parser.docx"
    doc = Document()
    doc.add_heading("Water", level=1)
    doc.add_paragraph("")
    doc.add_paragraph("Synthetic groundwater paragraph.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0,0).text = "Synthetic water table cell."
    doc.save(path)
    units, flags = extract(path)
    assert all(unit["page"] is None for unit in units)
    assert units[1]["paragraph"] == 3 and units[1]["section"] == "Water"
    assert units[-1]["locator"] == "docx-table:1:row:1:cell:1:paragraph:1"
    assert flags


def test_chunking_exact_text_and_offsets():
    text = "water baseline " * 100
    chunks = list(split_unit(text, 120))
    assert "".join(chunk[2] for chunk in chunks) == text
    assert all(text[start:end] == piece and len(piece) <= 120 for start,end,piece in chunks)


def test_idempotence_provenance_and_citations(tmp_path):
    root = tmp_path / "input"; root.mkdir()
    folder = root / "project"; folder.mkdir()
    make_pdf(folder / "parser.pdf")
    output = tmp_path / "output"
    catalog = tmp_path / "catalog.json"
    catalog.write_text(json.dumps({"project/parser.pdf": {
        "title": "Synthetic parser test", "title_reference": "test fixture definition",
        "source_type": "project", "licence": "test only"}}))
    first = ingest(root, output, catalog)
    snapshot = {p.name: p.read_bytes() for p in output.iterdir()}
    second = ingest(root, output, catalog)
    assert first == second
    assert snapshot == {p.name:p.read_bytes() for p in output.iterdir()}
    records = load_passages(output / "passages.jsonl")
    docs = {doc["document_id"]:doc for doc in json.loads((output/"documents.json").read_text())}
    for key, record in records.items():
        assert record["chunk_id"] == key
        assert record["document_id"] in docs
        assert record["original_filename"] == "parser.pdf"
        assert record["title"] == "Synthetic parser test"
        assert record["page"] == 1 and record["source_type"] == "project"
        assert record["publication_or_effective_date"] is None
        assert record["original_source_url"] is None
    with (output/"passages.csv").open(newline="") as stream:
        assert len(list(csv.DictReader(stream))) == len(records)
    assert retrieve_passages(records,"water baseline",source_type="project")
    assert not retrieve_passages(records,"unrelated volcano")
    assert not retrieve_passages(records,"water",source_type="scientific")


def test_duplicate_content_and_renames_do_not_duplicate_citations(tmp_path):
    root=tmp_path/"input"; root.mkdir()
    make_pdf(root/"a.pdf")
    (root/"b.pdf").write_bytes((root/"a.pdf").read_bytes())
    result=ingest(root,tmp_path/"out")
    assert result["documents"] == 1
    docs=json.loads((tmp_path/"out/documents.json").read_text())
    assert docs[0]["aliases"] == ["a.pdf","b.pdf"]


def test_unverified_metadata_is_not_silently_accepted():
    with pytest.raises(ValidationError):
        Metadata(publication_or_effective_date="2007")
    with pytest.raises(ValidationError):
        Metadata(issuing_body_or_author="Unverified")


def test_corrupt_file_is_reported(tmp_path):
    root=tmp_path/"in"; root.mkdir()
    (root/"bad.pdf").write_bytes(b"not a PDF")
    result=ingest(root,tmp_path/"out")
    assert result["passages"] == 0
    assert result["inventory"][0]["status"] == "failed"


def test_registration_is_only_a_dry_run(tmp_path):
    path=tmp_path/"passages.jsonl"; path.write_text('{"text":"synthetic test only"}\n')
    plan=registration_plan(path,"s3a://private-bucket/knowledge/passages.jsonl")
    assert plan["dry_run"] is True
    assert plan == registration_plan(path,"s3a://private-bucket/knowledge/passages.jsonl")
    assert "FORMAT JSON" in plan["statements"][1]
    with pytest.raises(ValueError):
        registration_plan(path,'s3a://bucket/evil" DROP TABLE x')


def test_live_query_requires_real_matching_passage(monkeypatch):
    import httpx
    chunk_id="chunk-"+"a"*64
    record={"chunk_id":chunk_id,"document_id":"doc-1","original_filename":"parser.pdf",
            "source_type":"project","title":"Parser fixture","page":1,"locator":"pdf-page:1","text":"Synthetic parser test"}
    def handle(req):
        sql=json.loads(req.content)["sql"]
        assert "SELECT" in sql and "LIMIT 1" in sql and chunk_id in sql
        return httpx.Response(200,json={"data":[list(record.values())]})
    original=httpx.Client
    monkeypatch.setattr(httpx,"Client",lambda **kw: original(transport=httpx.MockTransport(handle),**kw))
    config=Config("https://test/api/v1.0","secret","1")
    assert query_registered_passage(config,"knowledge.passages",{chunk_id:record},chunk_id)["verified"]
    wrong={**record,"text":"Different passage"}
    with pytest.raises(RuntimeError,match="did not match"):
        query_registered_passage(config,"knowledge.passages",{chunk_id:wrong},chunk_id)


def test_metadata_changes_preserve_citations_and_file_changes_invalidate_them(tmp_path):
    root=tmp_path/"input"; root.mkdir()
    path=root/"parser.pdf"; make_pdf(path)
    out=tmp_path/"output"
    ingest(root,out)
    before=set(load_passages(out/"passages.jsonl"))
    catalog=tmp_path/"metadata.json"
    catalog.write_text(json.dumps({"parser.pdf":{"title":"Parser title","title_reference":"test fixture"}}))
    ingest(root,out,catalog)
    assert set(load_passages(out/"passages.jsonl")) == before
    path.unlink()  # remove only this temporary test fixture to create a new revision
    make_pdf(path,"Changed synthetic parser-test water text.")
    ingest(root,out,catalog)
    assert not before.intersection(load_passages(out/"passages.jsonl"))
